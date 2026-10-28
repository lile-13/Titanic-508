"""从课程参考模板制作报告；使用 Codex 捆绑 Python 运行，随后由 Word 更新域。"""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
DEST=ROOT/"outputs/delivery_20261028_4690"
WORK=DEST/"work"
SOURCE=WORK/"reference_report.docx"
FINAL=DEST/"大作业综合报告-待填班级-待填学号-待填姓名.docx"


def draw_workflow():
    canvas=Image.new("RGB",(1860,390),"white");draw=ImageDraw.Draw(canvas)
    font=ImageFont.truetype("C:/Windows/Fonts/msyh.ttc",34)
    steps=[("原始数据","8693 / 4277"),("统一特征计数","公开无标签池"),("折内数据处理","填补与编码"),("分组验证冻结","参数与权重"),("模型与平台验证","重载及提交")]
    for i,lines in enumerate(steps):
        x=15+i*374
        draw.rectangle((x,55,x+330,290),outline="#202020",width=3)
        for j,text in enumerate(lines):draw.text((x+165,110+j*80),text,font=font,fill="black",anchor="mm")
        if i<4:
            draw.line((x+333,175,x+368,175),fill="black",width=4)
            draw.polygon([(x+368,175),(x+355,166),(x+355,184)],fill="black")
    canvas.save(ROOT/"reports/figures/course_workflow_20261028.png")


def main():
    WORK.mkdir(parents=True,exist_ok=True)
    draw_workflow()
    template=SOURCE if SOURCE.exists() else ROOT/"resources/course_report_template.docx"
    doc=Document(template)
    source_paragraphs=list(doc.paragraphs)
    hprops={level:deepcopy(source_paragraphs[index]._p.pPr) for level,index in [(1,28),(2,46),(3,47)]}
    cover=doc.tables[0]
    def cell_text(cell,text):
        p=cell.paragraphs[0]
        fmt=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        for child in list(p._p):
            if child.tag!=qn("w:pPr"):p._p.remove(child)
        r=p.add_run(text)
        if fmt is not None:r._r.insert(0,fmt)
        for old in list(cell._tc)[2:]:
            if old.tag==qn("w:p"):cell._tc.remove(old)
    cell_text(cover.cell(1,1),"Spaceship Titanic 乘客传送预测")
    selection=cover.cell(1,1).paragraphs[0]
    selection.style=doc.styles["Normal"]
    outline=OxmlElement("w:outlineLvl");outline.set(qn("w:val"),"9");selection._p.get_or_add_pPr().append(outline)
    for run in selection.runs:run.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    for row in range(6,10):
        for col in range(1,5):cell_text(cover.cell(row,col),"待填")
        cell_text(cover.cell(row,5),"")
    rubric=doc.tables[2]
    for row in range(2,7):cell_text(rubric.cell(row,3),"")
    body=doc._element.body
    start=source_paragraphs[28]._p
    boundary=source_paragraphs[96]._p
    removing=False
    for child in list(body):
        if child is start:removing=True
        if child is boundary:break
        if removing:body.remove(child)
    # Replace cached TOC with a fresh field. Native Word constructs the page numbers.
    for p in source_paragraphs[11:27]:
        if p._p.getparent() is body:body.remove(p._p)
    toc=doc.add_paragraph()
    begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin")
    instruction=OxmlElement("w:instrText");instruction.set(qn("xml:space"),"preserve")
    instruction.text=' TOC \\o "1-3" \\h \\z \\u '
    end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
    for e in [begin,instruction,end]:toc.add_run()._r.append(e)
    source_paragraphs[27]._p.addprevious(toc._p)

    def p(text="",level=0,newpage=False,center=False,bold=False,size=10.5,keep=False):
        item=doc.add_paragraph()
        boundary.addprevious(item._p)
        if level:
            item._p.insert(0,deepcopy(hprops[level]))
            item.paragraph_format.keep_with_next=True
            item.paragraph_format.first_line_indent=Pt(0)
        else:
            item.paragraph_format.line_spacing=1.5
            item.paragraph_format.first_line_indent=Pt(21 if not center else 0)
            item.paragraph_format.space_after=Pt(5)
        item.paragraph_format.widow_control=True
        if newpage:item.paragraph_format.page_break_before=True
        if center:item.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if keep:item.paragraph_format.keep_with_next=True
        r=item.add_run(text)
        r.font.name="Times New Roman"
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"黑体" if level else "宋体")
        r.font.size=Pt({1:15,2:14,3:12}.get(level,size))
        r.bold=bool(level or bold)
        r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        return item

    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers))
        boundary.addprevious(t._tbl)
        t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Inches(w)
        borders=OxmlElement("w:tblBorders")
        for edge in ["top","left","bottom","right","insideH","insideV"]:
            e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"000000");borders.append(e)
        t._tbl.tblPr.append(borders)
        margins=OxmlElement("w:tblCellMar")
        for edge in ["top","bottom","left","right"]:
            e=OxmlElement("w:"+edge);e.set(qn("w:w"),"85");e.set(qn("w:type"),"dxa");margins.append(e)
        t._tbl.tblPr.append(margins)
        for vals,row in [(headers,t.rows[0])]+[(vals,t.add_row()) for vals in rows]:
            no_split=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(no_split)
            for i,(c,text) in enumerate(zip(row.cells,vals)):
                c.width=Inches(widths[i]);c.vertical_alignment=1
                para=c.paragraphs[0];para.paragraph_format.line_spacing=1.2
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                r=para.add_run(str(text));r.font.name="Times New Roman";r.font.size=Pt(10)
                r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"宋体")
                r.bold=row is t.rows[0]  # explicitly set below as row proxies are reconstructed
        for r in t.rows[0].cells:
            for run in r.paragraphs[0].runs:run.bold=True
        repeat=OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(repeat)
        return t

    def image(path,caption,width=5.9):
        item=p(center=True,keep=True)
        item.add_run().add_picture(str(ROOT/path),width=Inches(width))
        p(caption,center=True,size=10)

    p("一．问题简介、目标",1)
    p("本项目围绕 Kaggle Spaceship Titanic 二分类竞赛，使用乘客的出发星球、冷冻睡眠、舱位、年龄和消费等信息，预测其是否被传送到另一维度（Transported）。截至2026年10月28日，已经完成数据清洗、特征工程、分组验证、模型搜索、集成和两次平台提交。当前已观测的最佳 Kaggle 得分为0.81108。")
    p("项目目标是建立从原始数据到提交文件的可复现流程，并以有效验证证据决定是否采用新方案。课程报告同时记录无收益实验、数据分布检查和评估局限，避免把参数搜索中的最高分当作稳定泛化表现。")
    p("二．小组成员工作",1)
    p("队伍状态：已组成四人队伍（组长于2026年10月28日确认）。Kaggle提交账户为Litianle4690，已提交文件名包含组长学号后四位4690。四人的姓名、完整学号、班级及贡献率由组长在提交前填写。")
    table(["角色","姓名","已交接的工作"],[
        ["组长","待填","CatBoost基线、统一训练入口、优化集成、平台提交及成果整合"],
        ["成员2","待填","数据探索、清洗策略、对抗验证和预处理模块"],
        ["成员3","待填","特征字典、CabinZone、分组消融及候选特征对照"],
        ["成员4","待填","随机森林与LightGBM对照、误差分析、概率融合及交接"]
    ],[.65,.65,5.0])
    p("后续整合与新增实验由组长统一组织，并使用代码与文档辅助工具完成。表中记录仓库交接分工；个人贡献率不由程序推断。")

    p("三．问题分析流程设计",1,newpage=True)
    p("流程首先固定原始数据与同行组折划分，再分析缺失、类别和消费分布。无标签的群组及姓氏计数统一在公开train与test特征池中计算；标签只进入各折模型训练。每折填补器、类别编码和缩放器仅拟合训练折，验证折复用。")
    image("reports/figures/course_workflow_20261028.png","图1 分析与实现流程",6.1)
    table(["阶段","输入与操作","产物与核验"],[
        ["数据检查","读取8693条训练记录和4277条测试记录；检查ID及标签","字段体检表、分布图、数据哈希"],
        ["特征工程","拆分同行组与舱位；统一计数；规则补缺","基础特征、消费特征、姓氏和位置候选"],
        ["分组验证","按GroupId固定五折；训练折内拟合预处理","逐折指标、OOF概率、错误分析"],
        ["模型选择","开发折搜索；冻结参数、权重和阈值","冻结配置及确认折记录"],
        ["训练交付","按冻结配置重训；模型重载；生成CSV","模型、4277行预测、哈希和平台截图"]
    ],[.8,2.8,2.7])
    p("验证设计的依据是数据中的同行关系。训练集共有6217个GroupId，测试集共有3063个，二者重叠为0；按组划分能够减少同组乘客分散到训练和验证两侧造成的过于乐观估计。姓氏可以跨组，也可以跨train和test，不能把姓氏直接等同于同行组。")
    p("模型开发时排除第5折，仅用第1至4折进行训练与选择，再按冻结方案评估第5折。但第5折曾在早期基线中被查看，因此本报告称其为历史确认折，不将其描述为从未查看的盲测。")

    p("四．问题分析实现步骤",1,newpage=True)
    p("4.1 数据集初步探索",2)
    p("1. 数据集获取",3)
    p("数据来自Kaggle官方Spaceship Titanic竞赛，原始文件为train.csv、test.csv与sample_submission.csv。训练集8693行、14列，测试集4277行、13列；差异为训练集包含目标Transported。PassengerId按字符串读取与对齐，保留组号和组内序号中的前导零。")
    p("2. 数据集基本情况",3)
    p("HomePlanet、Destination、CryoSleep、VIP及拆分后的Deck、Side为类别信息；Age与五项消费为数值信息。训练标签接近平衡，评价指标为Accuracy，因此以固定0.5阈值的准确率为主，同时保留Log loss衡量概率质量。")
    image("reports/figures/01_missing_rate.png","图2 原始字段缺失率",5.65)
    p("多个字段存在少量缺失，消费呈明显右偏且包含较多零值。清洗前保留缺失位置，以区分真实零消费和未知消费；异常值先结合业务规则检查，不直接删除高消费乘客。")

    p("4.2 数据预处理",2,newpage=True)
    p("规则补缺优先使用可核验的特征关系：已知CryoSleep为True时，五项缺失消费可填0；已知存在正消费时，缺失CryoSleep可填False。同组分类补缺仅传播已观察到唯一值的HomePlanet和Side，其他剩余缺失由训练折统计量或模型原生机制处理。")
    p("同组Destination与Deck可以不同，其中717个同行组的Destination不唯一；2217个姓氏中有1664个跨组。因此不能用组众数随意覆盖目的地或甲板，也不能把同姓视作同组。")
    table(["历史对照","准确率","解释"],[
        ["S0 不填充","0.8028","指定历史模型的参考结果"],
        ["S2 简单填充","0.8026","与不填充差异很小"],
        ["S3 增加缺失标记","与S2相同","未在该实验中获得收益"],
        ["S4 规则与组内填补","0.8044","小幅差异不足以证明稳定提高"]
    ],[2.0,.8,3.5])
    p("缺失标记与标签的最高相关性约0.025，只能说明本次线性关联较弱，不能据此证明缺失机制完全随机。因此项目把主要精力转向特征口径、模型验证与融合，而不是无限细化填补规则。")
    p("SurnameSize最初在train与test分别计算，因样本量不同形成来源指纹：两侧均值约5.33和3.51，且有1536个姓氏跨数据集出现。改为build_features_combined(train,test)统一计数后，成员2指定对抗验证AUC由0.7336降至0.5454。此统计不使用Transported，但依赖当前公开待预测批次，不能直接宣称适用于任意未来单行预测。")

    p("4.3 构建分析模型及预测",2,newpage=True)
    p("1. 模型构建",3)
    p("首先使用逻辑回归作为参考，再比较随机森林、LightGBM和CatBoost。统一one-hot基线为59列，其中CabinZone占6列；CatBoost原生类别基线为30列。编码不同意味着这些分数比较的是完整建模方案，而非在完全相同矩阵上的纯算法差异。")
    table(["模型与输入","五折准确率均值","用途"],[
        ["逻辑回归 59列","0.7954","线性参考"],
        ["随机森林 59列","0.7874","树模型对照"],
        ["LightGBM 59列","0.8108","成员4对照"],
        ["CatBoost 30列","0.8163","固定参数原始基线"],
        ["冻结集成","0.8157","经选择后的描述性五折结果"]
    ],[2.25,1.45,2.6])
    p("优化阶段完成CatBoost的32组和LightGBM的48组参数试验，并比较基础、消费、位置与姓名信息。选出两个CatBoost配置和一个LightGBM配置，以42、2026、103三个种子重新训练。融合权重在开发数据上冻结为0.25、0.25、0.50，最终分类阈值为0.5。")
    table(["成员模型","特征及预处理","轮数","权重"],[
        ["CatBoost","spend / native","555","0.25"],
        ["CatBoost","rich_names / native","748","0.25"],
        ["LightGBM","base / rule_group","327","0.50"]
    ],[1.4,2.9,.8,1.2])
    p("完整参数保存在frozen_plan.json。冻结后对五个训练分区分别训练，每个分区包含三个配置及三个种子，共保存45个模型。测试概率按配置权重、种子及五个分区平均。LightGBM这三个种子的预测相同，不将其误计为额外的独立多样性收益。")

    p("2. 模型评估",3,newpage=True)
    p("相同开发划分下，固定CatBoost准确率为0.811646，冻结集成为0.821999。历史确认折准确率由0.807825变为0.808976，净增加2个正确预测：修正32个旧错误，同时新增30个错误。按同行组配对bootstrap得到准确率差的95%区间约为[-0.0070,0.0101]，区间跨越0。")
    image("reports/figures/optimization_20261025_fold_comparison.png","图3 原始CatBoost与冻结集成的逐折准确率",5.9)
    p("最终五折集成均值0.815713，标准差0.011160；按样本加权OOF准确率0.815714，Log loss为0.378380。CatBoost基线均值0.816287，略高于集成。本地证据不足以证明集成稳定更好，后续平台对照提供了另一项观测。")
    p("保存模型重算的训练准确率均值为0.869033，验证准确率均值为0.815713，差距约0.05332。该检查用于发现过拟合迹象，没有据此回调已冻结参数。参与选择的开发折指标属于选择后描述性结果，不是新的无偏评估。")

    p("3. 模型选择与新增实验",3,newpage=True)
    p("2026年10月28日新增关系特征实验，使用按GroupId交叉拟合的姓氏目标均值，确保每名训练乘客及其整个同行组的标签均不参与自身编码。验证和测试阶段只使用训练部分的目标统计；未知姓氏回退训练先验。")
    p("该轮先比较8种配置，再对前两种配置补充两个种子，共12项配置与种子记录。最佳三种子融合开发准确率为0.820129，低于现有冻结集成的0.821999，且Log loss由0.379529变为0.380866，未达到预设门槛，因此没有生成新的提交文件。")
    diversity=ROOT/"reports/experiments/diversity_study_20261028/summary.json"
    if diversity.exists():
        result=json.loads(diversity.read_text(encoding="utf-8"));f=result["frozen"]
        p(f"随后比较支持向量机、随机树、直方图提升与不同CatBoost训练方式，共14种初始配置。最佳三种子方案的开发准确率为{f['development']['accuracy']:.6f}，现有集成为{f['incumbent_accuracy']:.6f}。实验状态："+("已生成待平台验证候选。" if result["status"]=="candidate_generated" else "未达到预设门槛，保留原集成。"))
    else:
        p("异构模型补充实验仍在运行，包括支持向量机、随机树、直方图提升和不同CatBoost训练方式。尚未完成的结果不纳入成绩结论；完整运行记录将保留在独立实验目录。")
    p("模型选择以完整流程为对象：新方案必须在固定开发口径下表现出足够收益，同时概率质量不明显恶化，才进入冻结重训。新增试验使用曾被查看的历史折，只作为继续开发的证据，不包装成全新的独立盲测。")
    full=ROOT/"reports/experiments/full_data_candidate_20261028/summary.json"
    if full.exists():
        result=json.loads(full.read_text(encoding="utf-8"))
        p(f"另生成一份训练方式对照：不修改原冻结参数，用全部8693条训练记录训练9个模型，再与原五折集成各占50%融合，得到submission_fullblend4690.csv。与原集成相比有{result['changed_test_labels']}个测试标签变化。此候选没有新的OOF分数，尚未提交Kaggle，不能宣称已经提高成绩；它独立保留，不替换已观测的0.81108方案。")
    p("4. 预测与提交",3)
    p("预测结果按sample_submission.csv的PassengerId顺序生成，包含PassengerId和Transported两列，共4277行。提交前核对ID唯一、顺序一致、无缺失，Transported为True或False。输出文件的SHA256保存在提交清单中；本地复制文件已核验一致，但平台截图不能证明服务器接收字节与本地文件哈希相同。")

    p("5. 运行与复现",3,newpage=True)
    p("Python依赖版本固定在requirements.txt。项目保存数据、折划分、核心源码和冻结配置哈希。对最终模型先执行序列化，再从文件重载并独立计算测试概率；比较浮点误差与最终分类，避免只验证内存中模型。")
    p("从源码重新训练历史冻结方案：以下命令使用全新的run-name，不重复参数搜索，也不将已有摘要误当作本机已经生成模型。")
    for text in ["python -m venv .venv",r".venv\Scripts\python -m pip install -r requirements.txt",r".venv\Scripts\python -m unittest discover -s tests -v",r".venv\Scripts\python -m scripts.refit_frozen --run-name my_refit --threads 4"]:
        item=p(text,size=9.5);item.paragraph_format.first_line_indent=Pt(0)
    p("若使用附带模型，可直接运行scripts.predict_ensemble并指定bundle和新的输出目录。复现命令检查原始数据及冻结源码；Git的LF/CRLF换行差异允许规范化，其他内容变化必须另开实验，不能冒充原方案复现。")
    repro=ROOT/"reports/checks/isolated_reproduction_20261028.json"
    if repro.exists() and json.loads(repro.read_text(encoding="utf-8")).get("passed"):
        p("已在不含模型权重或预测缓存的隔离源码目录完成冻结方案重训，并与原提交文件及测试概率比较通过。该验证使用同一台电脑和Python环境，证明了隔离目录复现；不将其写作其他组员或另一台机器的独立复现。")
    else:p("隔离源码目录重训的完整检查记录随源码包提供。尚未通过的项目不据此宣称已独立复现；以reports/checks中的实际结果为准。")
    p("交付中保留历史实验摘要、OOF预测、特征字典、Kaggle原始截图、冻结参数和最终预测。模型权重与本地输出被Git忽略，源码ZIP另附最终模型，避免误把GitHub已有源码等同于所有训练产物均已上传。")

    p("五．问题分析结果",1,newpage=True)
    p("平台先后完成两次提交：原CatBoost得分0.80313，冻结集成得分0.81108。集成提高0.00795，即0.795个百分点，成为本项目目前已观测的最佳平台方案。该分数来自组长提供的Complete截图，未记录名次，也不推断私有测试成绩或统计显著性。")
    image("reports/figures/kaggle_comparison_20261027.png","图4 Kaggle两次成功提交及得分对照",6.15)
    table(["平台提交文件","得分","状态"],[
        ["submission4690.csv","0.80313","Complete"],
        ["submission_ensemble4690.csv","0.81108","Complete"]
    ],[3.75,1.0,1.55])
    p("当前保留outputs/kaggle_ready_4690/submission_ensemble4690.csv作为已验证候选。后续生成的候选必须单独提交后才能记录真实榜分，不能用本地开发准确率替代Kaggle成绩。")
    p("组队截图待补：在此处插入Kaggle Team页面的四人队伍信息截图，并核对队名和成员显示名称满足课程后四位要求。四人姓名、完整学号、班级和贡献率由组长填写；不得以提交账户名推断这些信息。")

    p("结果讨论与后续改进",newpage=True,bold=True,size=12)
    p("本项目的主要改进包括统一跨数据集的姓氏计数口径、按同行组验证、在折内拟合数据处理、保存完整概率和模型，以及记录冻结方案与平台对照。对抗验证暴露的计数指纹比复杂缺失填补更值得优先修正。")
    p("局限包括样本规模较小、历史验证折被反复查看、开发期参数搜索引入选择偏差，以及平台公开分数只反映特定测试部分。训练与测试共享无标签统计是针对本竞赛公开批次的设计；应用于新批次时必须重新评估。")
    p("后续方向应优先寻找可解释且互补的预测信号，使用预先限定的实验规模和独立记录；只有当新方案在一致口径下获得足够证据，才替换现有候选。单纯增加迭代轮数或不断依据榜分修改阈值，不能保证提高真实泛化能力。")
    p("资料与产物索引",bold=True,size=12)
    for text in [
        "竞赛数据与评价：https://www.kaggle.com/competitions/spaceship-titanic",
        "CatBoost算法资料：https://catboost.ai/docs/",
        "LightGBM算法资料：https://lightgbm.readthedocs.io/",
        "分组验证及算法接口：https://scikit-learn.org/stable/",
        "复现实验：reports/experiments/optimization_20261025/；追加实验分别位于relational_study_20261028与diversity_study_20261028。",
        "成绩证据：reports/experiments/kaggle_submissions.json与reports/figures/kaggle_comparison_20261027.png。"
    ]:
        reference=p(text,size=10)
        reference.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p("本报告的实验与成绩记录截至2026年10月28日；封面日期沿用课程模板预印信息。课程提交前请补齐个人信息和Team截图，并更新目录。")

    staged=WORK/"staged_report.docx";doc.save(staged)
    # Preserve untouched source package entries byte-for-byte before Word refresh.
    replacements={"word/document.xml","word/_rels/document.xml.rels","[Content_Types].xml"}
    with zipfile.ZipFile(template) as original,zipfile.ZipFile(staged) as generated,zipfile.ZipFile(FINAL,"w",zipfile.ZIP_DEFLATED) as target:
        generated_names=set(generated.namelist());source_names=set(original.namelist())
        for name in original.namelist():target.writestr(name,generated.read(name) if name in replacements else original.read(name))
        for name in sorted(generated_names-source_names):target.writestr(name,generated.read(name))
    with zipfile.ZipFile(template) as original,zipfile.ZipFile(FINAL) as result:
        untouched=[name for name in original.namelist() if name not in replacements]
        assert all(original.read(name)==result.read(name) for name in untouched)
    (WORK/"package_preservation.json").write_text(json.dumps(dict(preserved_entries=len(untouched),
        changed_entries=sorted(replacements),source_sha256=hashlib.sha256(template.read_bytes()).hexdigest(),
        report_sha256_before_word=hashlib.sha256(FINAL.read_bytes()).hexdigest()),indent=2),encoding="utf-8")
    print(str(FINAL))


if __name__=="__main__":main()
